"""crRNA 特异性扫描（Cas12a2, RNA 靶向）

定位：Cas12a2 由靶 RNA 激活，"脱靶"= 转录组中与 spacer 部分错配、可能错误激活的
其他 RNA。本模块在转录本序列（FASTA，sense 链）中扫描 spacer 互补区的近似位点，
并检查位点 3' 侧的 PFS（默认取注册表条目 pfs 字段: 共识 GAAAG + 容忍错配数，
可用 --effector/--pfs/--pfs-tol 覆盖；PFS 位于 sense 链窗口 3' 下游，
即靶 RNA 上 spacer 5' 端侧翼，8D4A 实测排布）。

判定口径（与 PDB 8D4A 复合物一致）:
  - crRNA spacer 与靶 RNA 反向互补配对 → 扫描目标 = revcomp(spacer)
  - PFS 5'-GAAAG-3' 位于靶 RNA 互补区的 3' 下游（8D4A 靶 RNA 实测排布）
  - 错配数 <= --max-mismatch 的窗口计入位点表；0 错配+PFS = 预期靶点，
    1 错配位点尤其关键（对应 TP53 野生型 vs R248Q 突变型的单碱基鉴别场景）

用法:
  python crrna_specificity_scan.py --spacer <23nt> --fasta transcripts.fa
  python crrna_specificity_scan.py --spacer <23nt> --gene TP53   # 经 NCBI 拉 mRNA
输出: <out>.sites.csv + <out>.summary.json + 屏幕汇总
"""
import argparse
import csv
import json
import os
import sys
from urllib.parse import quote
from urllib.request import urlopen

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from scaffold_registry import resolve_pfs, pfs_mismatches  # noqa: E402

def _build_comp():
    table = str.maketrans({'A': 'T', 'C': 'G', 'G': 'C', 'T': 'A', 'U': 'A'})
    return table


COMP = _build_comp()


def revcomp(seq):
    return seq.translate(COMP)[::-1]


def normalize(seq):
    return str(seq).upper().replace('U', 'T')


def fetch_gene_seq(gene, *, organism='Homo sapiens', db='nuccore', timeout=30.0,
                   base_url='https://eutils.ncbi.nlm.nih.gov/entrez/eutils'):
    """从 NCBI 拉取基因 mRNA 序列（与旧链 design.py 同口径, 纯 urllib）。"""
    term = quote(f'{gene}[Gene] AND {organism}[Organism] AND mRNA[Filter]')
    url = f'{base_url}/esearch.fcgi?db={db}&term={term}&retmax=1&retmode=json'
    with urlopen(url, timeout=timeout) as resp:
        ids = json.loads(resp.read().decode('utf-8')).get('esearchresult', {}).get('idlist', [])
    if not ids:
        raise RuntimeError(f'NCBI 未找到基因 {gene} 的 mRNA 记录')
    url2 = f'{base_url}/efetch.fcgi?db={db}&id={ids[0]}&rettype=fasta&retmode=text'
    with urlopen(url2, timeout=timeout) as resp:
        lines = resp.read().decode('utf-8').splitlines()
    return ''.join(l.strip() for l in lines if not l.startswith('>'))


def read_fasta(path):
    """读多记录 FASTA, 返回 [(name, seq)]（DNA 字母表, U→T）。"""
    records, name, chunks = [], None, []
    with open(path, encoding='utf-8') as fh:
        for line in fh:
            line = line.strip()
            if line.startswith('>'):
                if name is not None:
                    records.append((name, normalize(''.join(chunks))))
                name, chunks = line[1:].split()[0] or f'record{len(records)}', []
            elif line:
                chunks.append(line)
    if name is not None:
        records.append((name, normalize(''.join(chunks))))
    if not records:
        raise ValueError(f'{path} 中没有 FASTA 记录')
    return records


def scan_record(name, seq, target, pfs_rule, max_mm, pfs_tol=0):
    """在单条转录本上找 target 的 <=max_mm 错配窗口, 并读 3' 下游 PFS。

    pfs_match = 与共识精确匹配(0 错配); pfs_tolerant = 错配数 <= pfs_tol
    （容忍模型, 容忍度默认取注册表 pfs.tolerant_mismatches, R5 统一口径）。"""
    L, lp = len(target), len(pfs_rule)
    sites = []
    for i in range(0, len(seq) - L - lp + 1):
        window = seq[i:i + L]
        mm = sum(a != b for a, b in zip(window, target))
        if mm <= max_mm:
            pfs = seq[i + L:i + L + lp]
            pmm = pfs_mismatches(pfs, pfs_rule)
            sites.append({'transcript': name, 'pos': i, 'mismatches': mm,
                          'window': window, 'pfs': pfs, 'pfs_mm': pmm,
                          'pfs_match': pmm == 0,
                          'pfs_tolerant': pmm is not None and pmm <= pfs_tol})
    return sites


def pfs_landscape_mode(out_prefix):
    """Scholz 2026 PFS 全枚举筛选(MOESM3)的实证景观 + 设计 PFS 校验(2026-09-09)。

    数据: 单 spacer x 全部 4^5=1024 种 5nt PFS 的 target/non-target depletion
    (Cas12a2 原生, 细胞内, GeCas12a2 RNP 口径)。
    输出: 位置碱基富集(top10% vs 全体) + 全量排序 + 四靶 panel 设计 PFS 的
    实证分位与鉴别裕量(mut vs wt depletion 比)。
    边界: 景观为 GeCas12a2/单一 spacer 上下文, 跨酶/跨 spacer 迁移未证;
    用作设计校验层而非硬过滤。
    """
    import json as _json

    import numpy as np
    import openpyxl
    src = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       '..', 'data', 'scholz2026_MOESM3.xlsx')
    wb = openpyxl.load_workbook(src, read_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))[1:]
    data = [(str(r[2]), float(r[3])) for r in rows if r[2] and r[3] is not None]
    if len(data) != 1024:
        raise SystemExit('PFS 景观异常: 期望 1024 行, 实得 %d' % len(data))
    dep = np.array([v for _, v in data])
    top_thr = float(np.percentile(dep, 90))
    top = [p for p, v in data if v >= top_thr]
    logo = {}
    for i in range(5):
        bases = [p[i] for p in top]
        logo['pos%d' % (i + 1)] = {b: round(bases.count(b) / len(bases), 3)
                                   for b in 'ACGU'}
    srt = sorted(data, key=lambda x: -x[1])
    rank_of = {p: i + 1 for i, (p, _) in enumerate(srt)}
    dep_of = dict(data)

    # 四靶 panel 设计 PFS 的实证校验(序列取自 data/agent/*.design.json,
    # R248Q 正典条目取 gRNA3_pick 口径)
    panel = {}
    agent_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                             '..', 'data', 'agent')
    mapping = {'TP53-R248Q': 'tp53_r248q.design.json',
               'KRAS-G12C': 'kras_g12c.design.json',
               'KRAS-G12D': 'kras_g12d.design.json',
               'TP53-R273H': 'tp53_r273h.design.json'}
    for tgt, fn in mapping.items():
        d = _json.load(open(os.path.join(agent_dir, fn), encoding='utf-8'))
        if tgt == 'TP53-R248Q':
            best = next((x for x in d['designs']
                         if x['spacer_dna'] == 'GTTCATGCCGCCCATGCAGGAACT'),
                        d['designs'][0])
        else:
            best = d['designs'][0]
        mut = best['pfs_mut'].replace('T', 'U')
        wt = (best.get('pfs_wt') or '').replace('T', 'U')
        entry = {'mut_pfs': mut, 'mut_rank': rank_of.get(mut),
                 'mut_depletion': dep_of.get(mut)}
        if wt and len(wt) == 5 and wt in rank_of:
            entry.update({'wt_pfs': wt, 'wt_rank': rank_of[wt],
                          'wt_depletion': dep_of[wt],
                          'discrimination_ratio': round(
                              dep_of[mut] / max(dep_of[wt], 1e-9), 2)})
        panel[tgt] = entry
        print('[%s] mut %s rank %s dep %s | wt %s rank %s dep %s | 比值 %s'
              % (tgt, mut, entry.get('mut_rank'), entry.get('mut_depletion'),
                 wt or '-', entry.get('wt_rank', '-'),
                 entry.get('wt_depletion', '-'),
                 entry.get('discrimination_ratio', '-')))
    payload = {
        'source': 'Scholz 2026 Nature s41586-026-10466-y MOESM3 '
                  '(PFS 全枚举筛选, 4^5=1024, Cas12a2 细胞内 depletion)',
        'boundary': 'GeCas12a2/单 spacer 上下文; 迁移未证, 作设计校验层非硬过滤',
        'depletion_stats': {'median': round(float(np.median(dep)), 3),
                            'p90_threshold': round(top_thr, 3),
                            'range': [round(float(dep.min()), 3),
                                      round(float(dep.max()), 3)]},
        'top10pct_base_enrichment': logo,
        'empirical_motif_reading': 'pos2-4 强 A 富集(0.57/0.76/0.60), '
                                   'pos1 C/G, pos5 混合——与 GAAAG 型 A-rich '
                                   'PFS 文献口径一致',
        'panel_design_check': panel,
        'panel_reading': '实证鉴别裕量(mut/wt depletion 比): R273H 4.68 最强'
                         '(mut 27.4%分位 vs wt 86.1%); R248Q 1.85 中等——mut 前 '
                         '9.5%(强激活)但 wt 侧 CGGAG 也中等激活(26.1%分位), '
                         'PFS 单独鉴别可能不足, 细胞层等位选择性主要依赖 protospacer '
                         '相同+PFS 差异的联合效应, 预注册 H1/H4 已覆盖该读数路径; '
                         'KRAS-G12C/G12D 设计 PFS 排 847/861 且比值<1(≈不激活, '
                         '红旗在案, 已退出本轮湿)。景观为 GeCas12a2/单 spacer 上下文, '
                         '跨酶迁移未证(Su 体系用 8D4A 口径独立设计)'}
    dst = out_prefix + '.pfs_landscape.json'
    _json.dump(payload, open(dst, 'w', encoding='utf-8'),
               ensure_ascii=False, indent=1)
    print('PFS 实证景观 -> %s' % dst)


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument('--spacer', default=None, help='固定 spacer(17-25nt ACGT/U)')
    ap.add_argument('--fasta', default=None, help='转录本 FASTA(可多条)')
    ap.add_argument('--gene', default=None, help='基因名(经 NCBI 拉 mRNA, 与 --fasta 二选一)')
    ap.add_argument('--effector', default='cas12a2',
                    help='注册表效应子条目(默认 cas12a2; PFS 规则默认取自该条目 pfs 字段)')
    ap.add_argument('--pfs', default=None,
                    help='PFS 共识序列(默认取注册表 pfs.consensus; 换体系用 --pfs 覆盖)')
    ap.add_argument('--pfs-tol', type=int, default=None,
                    help='PFS 容忍错配数(默认取注册表 pfs.tolerant_mismatches)')
    ap.add_argument('--max-mismatch', type=int, default=4)
    ap.add_argument('--pfs-landscape', action='store_true',
                    help='Scholz 2026 PFS 全枚举实证景观 + 四靶设计 PFS 校验')
    ap.add_argument('--out', default='crrna_specificity', help='输出前缀')
    args = ap.parse_args()

    if args.pfs_landscape:
        pfs_landscape_mode(args.out)
        return

    if not args.spacer:
        ap.error('需要 --spacer 或 --pfs-landscape')
    spacer = normalize(args.spacer)
    if not 17 <= len(spacer) <= 25 or set(spacer) - set('ACGT'):
        ap.error('--spacer 必须为 17-25nt ACGT/U')
    pfs_spec = resolve_pfs(args.effector, consensus=args.pfs,
                           tolerant_mismatches=args.pfs_tol)
    pfs_rule = pfs_spec['consensus']
    pfs_tol = pfs_spec['tolerant_mismatches']
    target = revcomp(spacer)  # spacer 与靶 RNA 反向互补

    if args.fasta:
        records = read_fasta(args.fasta)
    elif args.gene:
        seq = fetch_gene_seq(args.gene)
        records = [(args.gene, normalize(seq))]
    else:
        ap.error('需要 --fasta 或 --gene')

    sites = []
    for name, seq in records:
        sites.extend(scan_record(name, seq, target, pfs_rule, args.max_mismatch,
                                 pfs_tol))

    sites.sort(key=lambda s: (s['mismatches'], not s['pfs_match'], s['transcript'], s['pos']))
    by_mm = {}
    for s in sites:
        key = f"{s['mismatches']}mm" + ('+PFS' if s['pfs_match'] else '')
        by_mm[key] = by_mm.get(key, 0) + 1
    # 双模型脱靶计数(R5): exact=PFS 0 错配; tolerant=PFS 错配<=容忍度
    n_exact = sum(1 for s in sites if s['pfs_match'])
    n_tolerant = sum(1 for s in sites if s['pfs_tolerant'])

    print(f'spacer({len(spacer)}nt): {spacer}')
    print(f"扫描目标(revcomp): {target}  PFS: {pfs_rule}±{pfs_tol}"
          f"(位点 3' 下游, 来源 {pfs_spec['source']})")
    print(f'转录本: {len(records)} 条  命中位点(<= {args.max_mismatch} 错配): {len(sites)}')
    for key in sorted(by_mm):
        print(f'  {key}: {by_mm[key]}')
    print(f'  PFS 双模型计数: exact(0mm)={n_exact}  tolerant(<={pfs_tol}mm)={n_tolerant}')
    print('\n位点明细(错配升序):')
    print(f"{'转录本':<16}{'位置':>8}{'错配':>5}{'PFS':>8}{'窗口'}")
    for s in sites[:50]:
        print(f"{s['transcript']:<16}{s['pos']:>8}{s['mismatches']:>5}"
              f"{s['pfs'] + ('✓' if s['pfs_match'] else '✗'):>8}{s['window']}")

    out_csv = args.out + '.sites.csv'
    with open(out_csv, 'w', newline='', encoding='utf-8') as fh:
        writer = csv.DictWriter(fh, fieldnames=list(sites[0].keys()) if sites else
                                ['transcript', 'pos', 'mismatches', 'window', 'pfs', 'pfs_match'])
        writer.writeheader()
        writer.writerows(sites)
    out_json = args.out + '.summary.json'
    with open(out_json, 'w', encoding='utf-8') as fh:
        json.dump({'spacer_dna': spacer, 'target_scanned': target,
                   'pfs_rule': pfs_rule, 'pfs_tolerant_mismatches': pfs_tol,
                   'pfs_source': pfs_spec['source'],
                   'pfs_models_counts': {'exact_0mm': n_exact,
                                         f'tolerant_{pfs_tol}mm': n_tolerant},
                   'max_mismatch': args.max_mismatch, 'n_transcripts': len(records),
                   'n_sites': len(sites), 'by_mismatch': by_mm, 'sites': sites},
                  fh, ensure_ascii=False, indent=1)
    print(f'\n输出: {out_csv} / {out_json}')
    print('说明: 0 错配+PFS = 预期靶点; 1 错配位点对应单碱基鉴别压力(如 TP53 WT vs 突变);')
    print('      本扫描为纯序列近似评估, 激活与否还取决于实验条件, 最终特异性以湿实验为准。')


if __name__ == '__main__':
    main()
